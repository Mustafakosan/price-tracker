"""End-to-end tests: config -> checks -> history -> report -> alert."""

from decimal import Decimal
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from tracker.__main__ import main
from tracker.config import ConfigError, Product, load_settings
from tracker.core import run_checks
from tracker.notifier import TelegramNotifier, format_alert
from tracker.report import write_report
from tracker.storage import PriceStore

PAGE = '<div class="p"><span class="price">{}</span></div>'


@pytest.fixture
def store(tmp_path):
    with PriceStore(tmp_path / "prices.db") as s:
        yield s


def fake_shop(prices: dict[str, str]):
    """A fetch function that serves pages from a dict instead of the internet."""
    return lambda url: PAGE.format(prices[url])


PRODUCTS = [
    Product("Headphones", "https://shop.example/a", ".price", "TRY"),
    Product("Watch", "https://shop.example/b", ".price", "TRY"),
]


def test_first_run_sets_baseline_and_second_run_detects_change(store):
    first = run_checks(PRODUCTS, fake_shop({"https://shop.example/a": "1.299,90 TL", "https://shop.example/b": "2.499 TL"}), store)
    assert [r.previous for r in first] == [None, None]
    assert not any(r.is_alert(Decimal("0")) for r in first)

    second = run_checks(PRODUCTS, fake_shop({"https://shop.example/a": "1.149,90 TL", "https://shop.example/b": "2.499 TL"}), store)
    headphones, watch = second
    assert headphones.previous == Decimal("1299.90")
    assert headphones.current == Decimal("1149.90")
    assert headphones.change == Decimal("-150.00")
    assert headphones.change_pct == Decimal("-11.54")
    assert headphones.is_alert(Decimal("1"))
    assert watch.change == 0 and not watch.is_alert(Decimal("0"))
    assert len(store.history()) == 4


def test_min_change_threshold_filters_small_moves(store):
    run_checks(PRODUCTS[:1], fake_shop({"https://shop.example/a": "100,00 TL"}), store)
    (r,) = run_checks(PRODUCTS[:1], fake_shop({"https://shop.example/a": "100,50 TL"}), store)
    assert r.change_pct == Decimal("0.50")
    assert r.is_alert(Decimal("0.5"))
    assert not r.is_alert(Decimal("1"))


def test_one_broken_page_does_not_stop_the_run(store):
    def fetch(url):
        if url.endswith("/a"):
            raise OSError("connection reset")
        return PAGE.format("2.499 TL")

    a, b = run_checks(PRODUCTS, fetch, store)
    assert not a.ok and "connection reset" in a.error
    assert b.ok and b.current == Decimal("2499")
    assert len(store.history()) == 1  # failed checks are not stored


def test_report_has_summary_and_history(store, tmp_path):
    run_checks(PRODUCTS, fake_shop({"https://shop.example/a": "1.299,90 TL", "https://shop.example/b": "2.499 TL"}), store)
    results = run_checks(PRODUCTS, fake_shop({"https://shop.example/a": "1.149,90 TL", "https://shop.example/b": "2.699 TL"}), store)
    path = write_report(results, store.history(), tmp_path / "out" / "report.xlsx")

    wb = load_workbook(path)
    summary, history = wb["Summary"], wb["History"]
    assert summary["A2"].value == "Headphones"
    assert summary["C2"].value == pytest.approx(1149.90)
    assert summary["E2"].value == pytest.approx(-11.54)
    assert summary["A2"].fill.fgColor.rgb.endswith("E2F5E9")  # green: cheaper
    assert summary["A3"].fill.fgColor.rgb.endswith("FDE2E1")  # red: more expensive
    assert history.max_row == 5  # header + 4 checks


def test_alert_message_lists_changes_and_errors(store):
    run_checks(PRODUCTS[:1], fake_shop({"https://shop.example/a": "100 TL"}), store)
    (changed,) = run_checks(PRODUCTS[:1], fake_shop({"https://shop.example/a": "90 TL"}), store)
    broken = run_checks([Product("Broken <tag>", "https://shop.example/x", ".nope")], fake_shop({"https://shop.example/x": "1"}), store)

    msg = format_alert([changed], broken)
    assert "1 price change" in msg
    assert "100.00 TRY → <b>90.00 TRY</b> (-10.00%)" in msg
    assert "Broken &lt;tag&gt;" in msg  # product names are HTML-escaped


def test_telegram_failure_does_not_leak_token(caplog):
    import requests

    session = MagicMock()
    session.post.side_effect = requests.ConnectionError("https://api.telegram.org/botSECRET123/sendMessage")
    assert TelegramNotifier("SECRET123", "42", session=session).send("hi") is False
    assert "SECRET123" not in caplog.text


def test_config_validation(tmp_path):
    bad = tmp_path / "config.toml"
    bad.write_text('[[products]]\nname = "A"\nurl = "https://x"\n', encoding="utf-8")
    with pytest.raises(ConfigError, match="selector"):
        load_settings(bad)

    dup = tmp_path / "dup.toml"
    dup.write_text(
        '[[products]]\nname = "A"\nurl = "https://x"\nselector = ".p"\n'
        '[[products]]\nname = "A"\nurl = "https://y"\nselector = ".p"\n',
        encoding="utf-8",
    )
    with pytest.raises(ConfigError, match="Duplicate"):
        load_settings(dup)

    broken = tmp_path / "broken.toml"
    broken.write_text("products = [unclosed\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="not valid TOML"):
        load_settings(broken)


def test_dotenv_next_to_config_is_loaded_but_real_env_wins(tmp_path, monkeypatch):
    # setenv first so monkeypatch restores the original state after the test
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "placeholder")
    monkeypatch.delenv("TELEGRAM_BOT_TOKEN")
    monkeypatch.setenv("TELEGRAM_CHAT_ID", "from-env")
    (tmp_path / ".env").write_text("# comment\nTELEGRAM_BOT_TOKEN='abc:123'\nTELEGRAM_CHAT_ID=from-file\n", encoding="utf-8")
    config = tmp_path / "config.toml"
    config.write_text('[[products]]\nname = "A"\nurl = "https://x"\nselector = ".p"\n', encoding="utf-8")

    settings = load_settings(config)
    assert settings.telegram_token == "abc:123"
    assert settings.telegram_chat_id == "from-env"
    assert "abc:123" not in repr(settings)  # token hidden from logs/repr


def test_cli_end_to_end_with_local_pages(tmp_path, monkeypatch):
    monkeypatch.delenv("TELEGRAM_BOT_TOKEN", raising=False)
    page = tmp_path / "item.html"
    page.write_text(PAGE.format("£51.77"), encoding="utf-8")
    config = tmp_path / "config.toml"
    config.write_text(
        f'[scraping]\ndelay_seconds = 0\n[[products]]\nname = "Book"\nurl = "{page.as_uri()}"\nselector = ".price"\n',
        encoding="utf-8",
    )
    assert main(["--config", str(config)]) == 0
    page.write_text(PAGE.format("£45.00"), encoding="utf-8")
    assert main(["--config", str(config)]) == 0

    wb = load_workbook(tmp_path / "reports" / "price_report.xlsx")
    assert wb["Summary"]["B2"].value == pytest.approx(51.77)
    assert wb["Summary"]["C2"].value == pytest.approx(45.00)
    assert wb["Summary"]["F2"].value == "GBP"


def test_cli_returns_error_code_for_missing_config(tmp_path):
    assert main(["--config", str(tmp_path / "nope.toml")]) == 2

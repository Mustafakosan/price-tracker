"""Excel report: a colour-coded summary sheet plus the full price history."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import CheckResult
from .storage import PricePoint

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
UP_FILL = PatternFill("solid", fgColor="FDE2E1")      # price went up
DOWN_FILL = PatternFill("solid", fgColor="E2F5E9")    # price went down
ERROR_FILL = PatternFill("solid", fgColor="FFF4CE")
MONEY = "#,##0.00"
PERCENT = '+0.00"%";-0.00"%";0.00"%"'


def _header(ws, titles: list[str], widths: list[int]) -> None:
    ws.append(titles)
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _naive(dt):
    """Excel has no time zones; write UTC times without tzinfo."""
    return dt.replace(tzinfo=None)


def write_report(results: list[CheckResult], history: list[PricePoint], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Summary ---------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    _header(
        ws,
        ["Product", "Previous", "Current", "Change", "Change %", "Currency", "Checked (UTC)", "Status", "URL"],
        [30, 12, 12, 12, 11, 10, 18, 26, 50],
    )
    for r in results:
        status = "Error: " + r.error if r.error else ("New" if r.previous is None else "OK")
        ws.append([
            r.product,
            float(r.previous) if r.previous is not None else None,
            float(r.current) if r.current is not None else None,
            float(r.change) if r.change is not None else None,
            float(r.change_pct) if r.change_pct is not None else None,
            r.currency,
            _naive(r.checked_at),
            status,
            r.url,
        ])
        row = ws.max_row
        for col in (2, 3, 4):
            ws.cell(row=row, column=col).number_format = MONEY
        ws.cell(row=row, column=5).number_format = PERCENT
        ws.cell(row=row, column=7).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=row, column=9).hyperlink = r.url

        fill = None
        if r.error:
            fill = ERROR_FILL
        elif r.change and r.change > 0:
            fill = UP_FILL
        elif r.change and r.change < 0:
            fill = DOWN_FILL
        if fill:
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = fill
    ws.auto_filter.ref = ws.dimensions

    # --- History ---------------------------------------------------------
    hs = wb.create_sheet("History")
    _header(hs, ["Checked (UTC)", "Product", "Price", "Currency", "URL"], [18, 30, 12, 10, 50])
    for p in history:
        hs.append([_naive(p.checked_at), p.product, float(p.price), p.currency, p.url])
        hs.cell(row=hs.max_row, column=1).number_format = "yyyy-mm-dd hh:mm"
        hs.cell(row=hs.max_row, column=3).number_format = MONEY
    hs.auto_filter.ref = hs.dimensions

    wb.save(path)
    return path
